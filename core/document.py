from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def _replace_in_paragraph(para, mapping):
    """Substitui placeholders preservando formatação de cada run.

    Passo 1: substituição direta (preserva negrito/itálico/fonte 100%).
    Passo 2: mescla runs adjacentes quando o placeholder está dividido.
    """
    for run in para.runs:
        for key, val in mapping.items():
            if key in run.text:
                run.text = run.text.replace(key, val)

    for key, val in mapping.items():
        while True:
            texts = [r.text for r in para.runs]
            combined = "".join(texts)

            if key not in combined:
                break
            if any(key in t for t in texts):
                break  # já resolvido no passo 1

            start_pos = combined.index(key)
            end_pos = start_pos + len(key)
            pos = 0
            start_run = end_run = -1

            for i, t in enumerate(texts):
                if start_run == -1 and pos + len(t) > start_pos:
                    start_run = i
                if pos + len(t) >= end_pos:
                    end_run = i
                    break
                pos += len(t)

            if start_run == -1 or end_run == -1 or start_run == end_run:
                break

            merged = "".join(r.text for r in para.runs[start_run:end_run + 1])
            para.runs[start_run].text = merged.replace(key, val)
            for r in para.runs[start_run + 1:end_run + 1]:
                r.text = ""
            break


def _replace_in_element(element, mapping):
    for para_elem in element.iter(qn("w:p")):
        _replace_in_paragraph(Paragraph(para_elem, None), mapping)


def replace_placeholders(doc, mapping):
    """Substitui placeholders em todo o documento: corpo, tabelas, caixas de texto, cabeçalhos e rodapés."""
    for para in doc.paragraphs:
        _replace_in_paragraph(para, mapping)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_in_paragraph(para, mapping)

    for txbx in doc.element.body.iter(qn("w:txbxContent")):
        _replace_in_element(txbx, mapping)

    for section in doc.sections:
        for header_footer in [
            section.header, section.footer,
            section.even_page_header, section.even_page_footer,
            section.first_page_header, section.first_page_footer,
        ]:
            if header_footer is None:
                continue
            for para in header_footer.paragraphs:
                _replace_in_paragraph(para, mapping)
            for txbx in header_footer._element.iter(qn("w:txbxContent")):
                _replace_in_element(txbx, mapping)
