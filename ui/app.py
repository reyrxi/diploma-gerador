import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from docx import Document

from core.dates import format_date_full
from core.document import replace_placeholders
from ui.tabs import DisciplinasTab, build_tab_diploma, build_tab_geral, build_tab_historico


class DiplomaApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Gerador de Diploma e Historico")
        self.resizable(True, True)
        self.geometry("820x680")
        self._setup_style()
        self._build_ui()

    # -- Estilo ---------------------------------------------------------------

    def _setup_style(self):
        s = ttk.Style(self)
        s.theme_use("clam")
        s.configure("TFrame",        background="#f5f6fa")
        s.configure("TLabel",        background="#f5f6fa", font=("Segoe UI", 10))
        s.configure("TEntry",        font=("Segoe UI", 10))
        s.configure("TNotebook",     background="#e8eaf0")
        s.configure("TNotebook.Tab", font=("Segoe UI", 10, "bold"), padding=[12, 4])
        s.configure("Action.TButton", font=("Segoe UI", 11, "bold"),
                    padding=10, background="#1a5276", foreground="white")
        s.map("Action.TButton", background=[("active", "#1f618d")])
        self.configure(bg="#f5f6fa")

    # -- Layout ---------------------------------------------------------------

    def _build_ui(self):
        self._build_banner()
        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True, padx=10, pady=8)
        self.vars_geral = build_tab_geral(nb)
        self.vars_hist  = build_tab_historico(nb)
        self.vars_dip   = build_tab_diploma(nb)
        self.disc_tab   = DisciplinasTab(nb)
        self._build_action_bar()

    def _build_banner(self):
        banner = tk.Frame(self, bg="#154360", pady=12)
        banner.pack(fill="x")
        tk.Label(banner, text="Gerador de Diploma e Historico Escolar",
                 font=("Segoe UI", 15, "bold"), bg="#154360", fg="white").pack()
        tk.Label(banner, text="Preencha os dados e selecione o modelo Word para gerar o documento",
                 font=("Segoe UI", 9), bg="#154360", fg="#aed6f1").pack()

    def _build_action_bar(self):
        bar = tk.Frame(self, bg="#e8eaf0", pady=8)
        bar.pack(fill="x", side="bottom")
        ttk.Button(bar, text="Gerar Historico", style="Action.TButton",
                   command=self.gerar_historico).pack(side="left", padx=12)
        ttk.Button(bar, text="Gerar Diploma", style="Action.TButton",
                   command=self.gerar_diploma).pack(side="left", padx=4)
        ttk.Button(bar, text="Gerar Ambos", style="Action.TButton",
                   command=self.gerar_ambos).pack(side="left", padx=4)
        ttk.Button(bar, text="Limpar Campos",
                   command=self.limpar_campos).pack(side="right", padx=12)

    # -- Mapeamento de placeholders -------------------------------------------

    def _build_mapping(self) -> dict:
        g = self.vars_geral

        def get(key):
            return g[key].get().strip() if key in g else ""

        mapping = {
            "{{ALUNO}}":           get("aluno"),
            "{{DATA_NASC}}":       get("data_nasc"),
            "{{NACIONALIDADE}}":   get("nacionalidade"),
            "{{NATURALIDADE}}":    get("naturalidade"),
            "{{UF}}":              get("uf"),
            "{{FILIACAO_1}}":      get("filiacao_1"),
            "{{FILIACAO_2}}":      get("filiacao_2"),
            "{{CPF}}":             get("cpf"),
            "{{RG}}":              get("rg"),
            "{{ORGAO_EMISSOR}}":   get("orgao_emissor"),
            "{{CURSO_ANT}}":       get("curso_ant"),
            "{{ESTAB_ANT}}":       get("estab_ant"),
            "{{ANO_ANT}}":         get("ano_ant"),
            "{{CIDADE_ANT}}":      get("cidade_ant"),
            "{{TURMA}}":           get("turma"),
            "{{DATA_INICIO}}":     get("data_inicio"),
            "{{DATA_TERMINO}}":    get("data_termino"),
            "{{FREQUENCIA}}":      get("frequencia"),
            "{{RESULTADO}}":       get("resultado"),
            "{{COD_SISTEC}}":      get("cod_sistec"),
            "{{COD_CENSO}}":       get("cod_censo"),
            "{{CARGA_ESTAGIO}}":   get("carga_estagio"),
            "{{NATURALIDADE_UF}}": f"{get('naturalidade')}/{get('uf')}",
        }

        for key, var in self.disc_tab.get_nota_vars().items():
            mapping[f"{{{{{key.upper()}}}}}"] = var.get().strip()

        return mapping

    def _build_mapping_historico(self) -> dict:
        mapping = self._build_mapping()
        data_hist = self.vars_hist.get("data_hist", tk.StringVar()).get().strip()
        mapping["{{DATA_HIST}}"]       = format_date_full(data_hist)
        mapping["{{DATA_HIST_CURTA}}"] = data_hist
        return mapping

    def _build_mapping_diploma(self) -> dict:
        mapping = self._build_mapping()
        d = self.vars_dip

        def get(key):
            return d[key].get().strip() if key in d else ""

        data_dip = get("data_diploma")
        mapping.update({
            "{{NUM_REGISTRO}}":       get("num_registro"),
            "{{FOLHA}}":              get("folha"),
            "{{LIVRO}}":              get("livro"),
            "{{DATA_DIPLOMA}}":       format_date_full(data_dip),
            "{{DATA_DIPLOMA_CURTA}}": data_dip,
            "{{DATA_CONCLUSAO}}":     get("data_conclusao"),
            "{{NOME_ESCOLA}}":        get("nome_escola"),
            "{{MUNICIPIO_UF}}":       get("municipio_uf"),
            "{{EXPEDIDO_EM}}":        get("expedido_em"),
        })
        return mapping

    # -- Dialogos de arquivo --------------------------------------------------

    def _select_template(self, title):
        return filedialog.askopenfilename(
            title=f"Selecionar Modelo - {title}",
            filetypes=[("Documentos Word", "*.docx"), ("Todos", "*.*")],
        ) or None

    def _select_save(self, default_name):
        return filedialog.asksaveasfilename(
            title="Salvar como...",
            defaultextension=".docx",
            initialfile=default_name,
            filetypes=[("Documentos Word", "*.docx")],
        ) or None

    # -- Geracao de documentos ------------------------------------------------

    def gerar_historico(self):
        tpl = self._select_template("Historico Escolar")
        if not tpl:
            return
        aluno = self.vars_geral.get("aluno", tk.StringVar()).get().strip()
        out = self._select_save(f"Historico_{aluno or 'aluno'}.docx")
        if not out:
            return
        try:
            doc = Document(tpl)
            replace_placeholders(doc, self._build_mapping_historico())
            doc.save(out)
            messagebox.showinfo("Sucesso", f"Historico gerado!\n{out}")
        except Exception as e:
            messagebox.showerror("Erro", str(e))

    def gerar_diploma(self):
        tpl = self._select_template("Diploma")
        if not tpl:
            return
        aluno = self.vars_geral.get("aluno", tk.StringVar()).get().strip()
        out = self._select_save(f"Diploma_{aluno or 'aluno'}.docx")
        if not out:
            return
        try:
            doc = Document(tpl)
            replace_placeholders(doc, self._build_mapping_diploma())
            doc.save(out)
            messagebox.showinfo("Sucesso", f"Diploma gerado!\n{out}")
        except Exception as e:
            messagebox.showerror("Erro", str(e))

    def gerar_ambos(self):
        self.gerar_historico()
        self.gerar_diploma()

    def limpar_campos(self):
        for d in [self.vars_geral, self.vars_hist, self.vars_dip]:
            for var in d.values():
                var.set("")
        self.disc_tab.clear()
